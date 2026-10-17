"""Read a deck through PowerPoint: slide stills, speaker notes, timings and fingerprints.

The source deck is never opened directly. It is copied to ``<output>/.work/deck.pptx`` first, so
the deck can stay open in PowerPoint while it is being reviewed.
"""

import hashlib
import json
import logging
import shutil
from dataclasses import asdict, dataclass
from pathlib import Path

import comtypes.client
from pptx import Presentation

log = logging.getLogger("presgen")

MSO_MEDIA = 16                 # MsoShapeType.msoMedia
EFFECT_MEDIA_PLAY = 83         # MsoAnimEffect.msoAnimEffectMediaPlay
TRIGGER_WITH_PREVIOUS = 2      # MsoAnimTriggerType
STILL_SIZE = (1920, 1080)


@dataclass
class SlideInfo:
    number: int                # PowerPoint slide number (hidden slides keep their number)
    slide_id: int              # PowerPoint SlideID: stable when slides are moved or inserted
    hidden: bool
    title: str
    notes: str
    fingerprint: str = ""      # changes whenever what the viewer sees or hears changes
    anim_secs: float = 0.0     # when the last animation (or embedded media) finishes
    planned_secs: float = 0.0  # the slide's own auto-advance time, if it has one

    @property
    def stem(self) -> str:
        return f"{self.number:02d}"


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 16), b""):
            h.update(block)
    return h.hexdigest()


class PowerPoint:
    """One PowerPoint COM session. Quits PowerPoint afterwards only if nothing else is open in it."""

    def __enter__(self):
        self.app = comtypes.client.CreateObject("PowerPoint.Application")
        return self

    def open(self, path: Path):
        # Read-only and windowless: the work copy is changed in memory (hidden flags, timings) only.
        return self.app.Presentations.Open(str(path.resolve()), ReadOnly=True, Untitled=True, WithWindow=False)

    def __exit__(self, *exc):
        try:
            if self.app.Presentations.Count == 0:
                self.app.Quit()
        except Exception as e:  # PowerPoint may already have gone
            log.debug(f"PowerPoint quit: {e}")


def work_copy(pptx: Path, out: Path) -> Path:
    work = out / ".work"
    work.mkdir(parents=True, exist_ok=True)
    dst = work / "deck.pptx"
    shutil.copy2(pptx, dst)
    return dst


def _anim_end(slide) -> float:
    """Simulate the slide's main animation sequence and return when its last effect ends."""
    seq = slide.TimeLine.MainSequence
    group_start = group_end = 0.0
    for j in range(1, seq.Count + 1):
        effect = seq(j)
        timing = effect.Timing
        if j > 1 and timing.TriggerType != TRIGGER_WITH_PREVIOUS:
            group_start = group_end  # "After previous" / "On click" start a new group
        dur = timing.Duration
        if effect.EffectType == EFFECT_MEDIA_PLAY:
            try:
                dur = effect.Shape.MediaFormat.Length / 1000
            except Exception:
                pass
        group_end = max(group_end, group_start + timing.TriggerDelayTime + dur)
    return group_end


def _timing_signature(slide) -> list:
    seq = slide.TimeLine.MainSequence
    sig = []
    for j in range(1, seq.Count + 1):
        e = seq(j)
        sig.append([e.EffectType, e.Timing.TriggerType, round(e.Timing.TriggerDelayTime, 2),
                    round(e.Timing.Duration, 2), e.Shape.Name])
    return sig


def _media_secs(slide) -> float:
    secs = 0.0
    for shape in slide.Shapes:
        if shape.Type == MSO_MEDIA:
            try:
                secs = max(secs, shape.MediaFormat.Length / 1000)
            except Exception:
                pass
    return secs


def _write_if_changed(path: Path, text: str) -> None:
    """Keep the file (and its timestamp, which the audio cache relies on) when the text is unchanged."""
    if path.exists() and path.read_text(encoding="utf-8") == text:
        return
    path.write_text(text, encoding="utf-8")


def _export_png(slide, path: Path) -> None:
    """Export a still, replacing the existing file only if the image changed (keeps its timestamp)."""
    tmp = path.with_suffix(".tmp.png")
    slide.Export(str(tmp.resolve()), "PNG", *STILL_SIZE)
    if path.exists() and sha256_file(tmp) == sha256_file(path):
        tmp.unlink()
    else:
        tmp.replace(path)


def inspect_deck(pptx: Path, out: Path) -> list:
    """Export stills and notes for every visible slide, and fingerprint each one."""
    out.mkdir(parents=True, exist_ok=True)
    work = work_copy(pptx, out)

    prs = Presentation(str(work))
    by_id = {}
    for s in prs.slides:
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        title = s.shapes.title.text if s.shapes.title is not None and s.shapes.title.has_text_frame else ""
        media = sorted(sha256_bytes(rel.target_part.blob) for rel in s.part.rels.values()
                       if not rel.is_external and rel.reltype.rsplit("/", 1)[-1] in ("image", "media", "video", "audio"))
        by_id[s.slide_id] = (notes, title.strip(), media)

    slides = []
    log.info(f"Opening {pptx.name} in PowerPoint")
    with PowerPoint() as pp:
        pres = pp.open(work)
        try:
            for i in range(1, pres.Slides.Count + 1):
                slide = pres.Slides(i)
                notes, title, media = by_id.get(slide.SlideID, ("", "", []))
                info = SlideInfo(number=i, slide_id=slide.SlideID, hidden=bool(slide.SlideShowTransition.Hidden),
                                 title=title or f"Slide {i}", notes=notes)
                if info.hidden:
                    log.info(f"Slide {i} is hidden: skipped")
                    slides.append(info)
                    continue
                png = out / f"slide_{info.stem}.png"
                _export_png(slide, png)
                _write_if_changed(out / f"text_{info.stem}.txt", notes)
                info.anim_secs = round(max(_anim_end(slide), _media_secs(slide)), 2)
                t = slide.SlideShowTransition
                info.planned_secs = round(t.AdvanceTime, 2) if t.AdvanceOnTime else 0.0
                info.fingerprint = sha256_bytes(json.dumps({
                    "notes": notes, "still": sha256_file(png), "media": media,
                    "timing": _timing_signature(slide),
                }, sort_keys=True).encode())[:16]
                slides.append(info)
        finally:
            pres.Close()

    (out / "slides.json").write_text(json.dumps([asdict(s) for s in slides], indent=2), encoding="utf-8")
    visible = [s for s in slides if not s.hidden]
    log.info(f"Inspected {len(slides)} slides ({len(visible)} visible)")
    return slides


def load_slides(out: Path) -> list:
    """The slide list from the last inspection, without opening PowerPoint."""
    path = out / "slides.json"
    if not path.exists():
        return []
    return [SlideInfo(**d) for d in json.loads(path.read_text(encoding="utf-8"))]
